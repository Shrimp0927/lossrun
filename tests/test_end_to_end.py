"""Runs the command on the sample PDFs and compares with tests/expected_report.json."""

import json
from pathlib import Path

from helpers import SAMPLES
from openpyxl import load_workbook
from typer.testing import CliRunner

from lossrun.cli import app

EXPECTED = Path(__file__).parent / "expected_report.json"
SAMPLE_FILES = [str(path) for path in sorted(SAMPLES.glob("*.pdf"))]
OPTIONS = ["--as-of", "2026-10-02", "--lines", "GL,WC,auto,property", "--years", "5"]

runner = CliRunner()


def test_samples_produce_the_expected_report(tmp_path):
    out = tmp_path / "report.json"
    result = runner.invoke(app, ["process", *SAMPLE_FILES, *OPTIONS, "--json", str(out)])
    assert result.exit_code == 0, result.output
    assert json.loads(out.read_text()) == json.loads(EXPECTED.read_text())


def test_samples_produce_a_workbook_with_one_row_per_record(tmp_path):
    out = tmp_path / "report.xlsx"
    runner.invoke(app, ["process", *SAMPLE_FILES, *OPTIONS, "--out", str(out)])
    expected = json.loads(EXPECTED.read_text())
    rows = {sheet.title: sheet.max_row - 1 for sheet in load_workbook(out)}
    assert rows == {
        "Claims": len(expected["claims"]),
        "Summary": len(expected["summary"]),
        "Flags": len(expected["flags"]),
        "Sources": len(expected["files"]),
    }


def test_file_that_cannot_be_extracted_exits_with_code_1(tmp_path):
    missing = tmp_path / "missing.pdf"
    result = runner.invoke(app, ["process", str(missing), SAMPLE_FILES[0]])
    assert (result.exit_code, "FAILED  missing.pdf" in result.output) == (1, True)


def test_unknown_line_of_coverage_is_rejected():
    result = runner.invoke(app, ["process", SAMPLE_FILES[0], "--lines", "GL,marine"])
    assert result.exit_code == 2
