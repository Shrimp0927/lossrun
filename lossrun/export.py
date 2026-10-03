"""Write a Report as an Excel workbook or JSON."""

from datetime import date
from decimal import Decimal
from enum import Enum
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel

from lossrun.models import Claim, Flag, PolicySummary, Report, SourceInfo

MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "yyyy-mm-dd"
HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MAX_WIDTH = 90


def write_json(report: Report, path: Path) -> None:
    path.write_text(report.model_dump_json(indent=2) + "\n")


def write_xlsx(report: Report, path: Path) -> None:
    workbook = Workbook()
    fill_sheet(workbook.active, "Claims", Claim, report.claims)
    fill_sheet(workbook.create_sheet(), "Summary", PolicySummary, report.summary)
    fill_sheet(workbook.create_sheet(), "Flags", Flag, report.flags)
    fill_sheet(workbook.create_sheet(), "Sources", SourceInfo, report.files)
    workbook.save(path)


def cell_value(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return value


def fill_sheet(sheet: Worksheet, title: str, model: type[BaseModel], rows: list) -> None:
    """Field names as headers, one row per object."""
    sheet.title = title
    fields = list(model.model_fields)
    sheet.append(fields)
    for cell in sheet[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL

    widths = [len(field) for field in fields]
    for row in rows:
        values = [getattr(row, field) for field in fields]
        sheet.append([cell_value(v) for v in values])
        for i, (cell, value) in enumerate(zip(sheet[sheet.max_row], values)):
            widths[i] = max(widths[i], len(str(cell.value or "")))
            if isinstance(value, Decimal):
                cell.number_format = MONEY_FORMAT
            elif isinstance(value, date):
                cell.number_format = DATE_FORMAT
            elif isinstance(value, str) and len(value) > MAX_WIDTH:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = min(width + 3, MAX_WIDTH)
